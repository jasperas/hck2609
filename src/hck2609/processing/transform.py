import json
from email import policy
from email.parser import BytesParser
from pathlib import Path

from pypdf import PdfReader

from hck2609.contracts import CleanData, RawData

import json
from pptx import Presentation
import pandas as pd
from docx import Document

def pptx_to_json(file_path):
    """
    Reads a PowerPoint file, extracts metadata, joins slide body elements 
    with a newline, and formats the output into the desired JSON structure.
    """
    prs = Presentation(file_path)
    presentation_data = []
    
    # Extract presentation-level core properties for metadata
    core_props = prs.core_properties
    doc_title = core_props.title if core_props.title else file_path
    doc_author = core_props.author if core_props.author else "Unknown"
    
    # Iterate through each slide in the presentation
    for slide_index, slide in enumerate(prs.slides):
        body_elements = []
        
        # Extract text from shapes on the slide
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        body_elements.append(text)
        
        # Join the text elements with a newline character
        combined_body = "\n".join(body_elements)
        
        # Construct the structured dictionary for the slide
        slide_entry = {
            "DocType": "ppt",
            "metadata": {
                "presentation_title": doc_title,
                "author": doc_author,
                "slide_number": slide_index + 1
            },
            "body": combined_body
        }
        
        presentation_data.append(slide_entry)
        
    # Convert the Python structure to a formatted JSON string
    return json.dumps(presentation_data, indent=4)

def docx_to_json(file_path):
    """
    Reads a Word document (.docx), extracts its paragraphs and table content for the body,
    pulls metadata from core properties, and formats it into the JSON structure.
    """
    doc = Document(file_path)
    
    # Extract metadata from core properties
    core_props = doc.core_properties
    metadata = {
        "title": core_props.title if core_props.title else "",
        "author": core_props.author if core_props.author else "",
        "last_modified_by": core_props.last_modified_by if core_props.last_modified_by else "",
        "revision": core_props.revision if core_props.revision else 1
    }
    
    body_elements = []
    
    # Extract text from paragraphs
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            body_elements.append(text)
            
    # Extract text from tables if the document contains any
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    body_elements.append(text)
                    
    # Join all body elements with a newline character
    combined_body = "\n".join(body_elements)
    
    # Construct the final JSON dictionary structure
    document_data = {
        "DocType": "word",
        "metadata": metadata,
        "body": combined_body
    }
    
    # Convert to a formatted JSON string
    return json.dumps(document_data, indent=4)

def excel_to_json(file_path):
    """
    Reads an Excel file, iterates through each sheet, extracts metadata,
    converts sheet rows into a newline-separated body string, and formats into JSON.
    """
    xls = pd.ExcelFile(file_path)
    presentation_data = []
    
    # Iterate through all sheets in the Excel workbook
    for sheet_name in xls.sheet_names:
        df = pd.read_excel(file_path, sheet_name=sheet_name)
        
        body_elements = []
        
        # Include the sheet name as context in the body or header
        body_elements.append(f"Sheet: {sheet_name}")
        
        # Convert dataframe rows into string representations
        for _, row in df.iterrows():
            # Drop NaN values and format row elements
            row_vals = [str(val).strip() for val in row.values if pd.notna(val) and str(val).strip() != ""]
            if row_vals:
                body_elements.append(" | ".join(row_vals))
                
        # Join all body elements with a newline character
        combined_body = "\n".join(body_elements)
        
        # Construct the structured dictionary for the sheet
        sheet_entry = {
            "DocType": "excel",
            "metadata": {
                "file_name": file_path,
                "sheet_name": sheet_name,
                "total_rows": len(df)
            },
            "body": combined_body
        }
        
        presentation_data.append(sheet_entry)
        
    # Convert the Python structure to a formatted JSON string
    return json.dumps(presentation_data, indent=4)


def clean(raw: RawData) -> CleanData:
    """Clean and normalise raw data. Stub: strips names, drops missing rows."""
    df = raw.dropna().copy()
    df["name"] = df["name"].str.strip().str.lower()
    return df


# def clean(raw: RawData) -> CleanData:
#     """Clean and normalise raw data. Stub: strips names, drops missing rows."""
#     df = raw.dropna().copy()
#     df["name"] = df["name"].str.strip().str.lower()
#     return df


def email_to_json(path: Path) -> dict:
    """Convert an .eml file to {"type", "metadata", "body"}."""
    msg = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
    body = msg.get_body(preferencelist=("plain", "html"))
    return {
        "type": "email",
        "metadata": {
            "from": msg["From"],
            "to": msg["To"],
            "cc": msg["Cc"],
            "subject": msg["Subject"],
            "date": msg["Date"],
            "message_id": msg["Message-ID"],
            "attachments": [a.get_filename() for a in msg.iter_attachments()],
        },
        "body": body.get_content() if body else "",
    }


def sharepoint_to_json(path: Path) -> dict:
    """Convert a SharePoint item export to {"type", "metadata", "body"}."""
    item = json.loads(path.read_text(encoding="utf-8"))
    body = item.pop("Content", "")
    return {"type": "sharepoint", "metadata": item, "body": body}


def teams_to_json(path: Path) -> dict:
    """Convert a Teams channel thread to {"type", "metadata", "body"}.

    The body is the list of messages, so replies and reactions stay available.
    """
    thread = json.loads(path.read_text(encoding="utf-8"))
    messages = thread.pop("messages", [])
    return {"type": "teams_thread", "metadata": thread, "body": messages}


def pdf_to_json(path: Path) -> dict:
    """Convert a PDF to {"type", "metadata", "body"} (text layer only, no OCR)."""
    reader = PdfReader(path)
    info = {k.lstrip("/"): v for k, v in (reader.metadata or {}).items()}
    return {
        "type": "pdf",
        "metadata": {**info, "pages": len(reader.pages)},
        "body": "\n\n".join(page.extract_text() for page in reader.pages),
    }


CONVERTERS = {
    "01_emails": email_to_json,
    "02_sharepoint": sharepoint_to_json,
    "03_teams_messages": teams_to_json,
    "04_pdf": pdf_to_json,
}


# def convert_raw(
#     raw_dir: Path = Path("data/raw"), out_dir: Path = Path("data/processed")
# ) -> list[Path]:
#     """Convert folders 01-04 of raw_dir to flat JSON files in out_dir."""
#     out_dir.mkdir(parents=True, exist_ok=True)
#     written = []
#     for folder, convert in CONVERTERS.items():
#         for path in sorted(p for p in (raw_dir / folder).iterdir() if p.is_file()):
#             out = out_dir / f"{path.stem}.json"
#             out.write_text(
#                 json.dumps(convert(path), indent=2, ensure_ascii=False, default=str),
#                 encoding="utf-8",
#             )
#             written.append(out)
#     return written
